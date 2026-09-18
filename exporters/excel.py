'''店長確認用 Excel の生成。openpyxl による装飾はここに閉じ込める。'''
import io

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.shift_rules import is_day_off

# 左4列（従業員コード〜希望出勤時間）を固定する
FREEZE_PANES = 'E2'

MIN_COLUMN_WIDTH = 12
COLUMN_WIDTH_FACTOR = 2

FILL_HEADER = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')  # グレー
FILL_SATURDAY = PatternFill(start_color='E6F2FF', end_color='E6F2FF', fill_type='solid')  # 薄い青
FILL_SUNDAY = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')  # 薄い赤

FONT_HEADER = Font(name='メイリオ', size=10, bold=True)
FONT_NORMAL = Font(name='メイリオ', size=10)
FONT_OFF = Font(name='メイリオ', size=10, bold=True, color='FF0000')  # 赤字・太字

BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center')


def _is_spacer_row(row) -> bool:
    '''部門の区切りとして挿入した空白行かどうか（従業員コードも名前も空）。'''
    if row[0].row <= 1:
        return False
    val_code = str(row[0].value or '').strip()
    val_name = str(row[1].value or '').strip()
    return val_code == '' and val_name == ''


def _is_day_column(col_name: str) -> bool:
    '''「1日（水）」のような日付列かどうか'''
    return '（' in col_name


def _style_worksheet(ws) -> None:
    ws.freeze_panes = FREEZE_PANES

    header = {
        cell.column: str(cell.value or '')
        for cell in ws[1]
    }

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        if _is_spacer_row(row):
            continue

        for cell in row:
            cell.border = BORDER_THIN
            col_name = header.get(cell.column, '')

            if cell.row == 1:
                cell.fill = FILL_HEADER
                cell.font = FONT_HEADER
                cell.alignment = ALIGN_CENTER
                continue

            cell.font = FONT_NORMAL
            cell.alignment = ALIGN_CENTER

            # 土曜と日曜を塗り分ける
            if '（土）' in col_name:
                cell.fill = FILL_SATURDAY
            elif '（日）' in col_name:
                cell.fill = FILL_SUNDAY

            # 日付列で「休」「有」なら赤字にする
            if _is_day_column(col_name) and cell.value and is_day_off(cell.value):
                cell.font = FONT_OFF

    # 列幅を見やすく自動調整
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len * COLUMN_WIDTH_FACTOR, MIN_COLUMN_WIDTH)


def build_shift_workbook(df: pd.DataFrame, sheet_name: str) -> bytes:
    '''整形済みの DataFrame から、装飾付き Excel をメモリ上で生成して bytes で返す。'''
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        _style_worksheet(writer.sheets[sheet_name])
    return output.getvalue()
