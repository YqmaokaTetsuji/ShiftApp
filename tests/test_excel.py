import io

import pandas as pd
from openpyxl import load_workbook

from exporters.excel import build_shift_workbook


def _frame():
    return pd.DataFrame(
        [
            {'従業員コード': '3', '名前': 'C', '部門': '季節AV',
             '希望出勤時間': 80, '3日（土）': '休', '4日（日）': '出'},
            {col: '' for col in
             ['従業員コード', '名前', '部門', '希望出勤時間', '3日（土）', '4日（日）']},
            {'従業員コード': '5', '名前': 'A', '部門': '家電',
             '希望出勤時間': 60, '3日（土）': '有', '4日（日）': '11-19'},
        ]
    )


def _load(df):
    data = build_shift_workbook(df, 'テスト')
    return load_workbook(io.BytesIO(data))['テスト']


def test_workbook_contains_sheet_and_data():
    ws = _load(_frame())
    assert ws.max_row == 4  # 見出し + 3行
    assert ws['A1'].value == '従業員コード'
    assert ws['B2'].value == 'C'


def test_header_is_styled_and_panes_frozen():
    ws = _load(_frame())
    assert ws.freeze_panes == 'E2'
    assert ws['A1'].font.bold is True
    assert ws['A1'].fill.start_color.rgb.endswith('D9D9D9')


def test_weekend_columns_are_filled():
    ws = _load(_frame())
    # E列 = 3日（土）, F列 = 4日（日）
    assert ws['E2'].fill.start_color.rgb.endswith('E6F2FF')
    assert ws['F2'].fill.start_color.rgb.endswith('FFE6E6')


def test_day_off_cells_are_red_and_bold():
    ws = _load(_frame())
    assert ws['E2'].value == '休'
    assert ws['E2'].font.color.rgb.endswith('FF0000')
    assert ws['E4'].value == '有'
    assert ws['E4'].font.color.rgb.endswith('FF0000')
    # 出勤日は赤くしない
    assert ws['F2'].value == '出'
    assert ws['F2'].font.bold is False


def test_spacer_row_is_left_unstyled():
    ws = _load(_frame())
    # 3行目は部門の区切りとして挿入した空白行
    assert ws['A3'].value in (None, '')
    assert ws['A3'].border.left.style is None


def test_column_widths_have_minimum():
    ws = _load(_frame())
    assert all(dim.width >= 12 for dim in ws.column_dimensions.values())
